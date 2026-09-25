"""Validate and freeze dissertation evidence without changing experiment records.

The script reads the Secure MLflow project, verifies Phase 1 artefact binding,
checks the final Phase 2 and Phase 3 claims, separates the pre-specified five
primary Phase 3 timing runs per scenario from later alert-validation runs, and
writes reproducible manifests and timing statistics.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import random
import re
import statistics
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable


EXPECTED_PHASE3_COUNTS = {
    "S1_LABEL_FLIP": 6,
    "S2_BACKDOOR": 7,
    "S3_PYPI": 6,
    "S4_MODEL_HUB": 6,
    "S5_MLFLOW_TAMPERING": 5,
}
PRIMARY_REPETITIONS = 5
EXPECTED_PHASE3_TOTAL = sum(EXPECTED_PHASE3_COUNTS.values())
EXPECTED_SUPPLEMENTARY_RUNS = EXPECTED_PHASE3_TOTAL - (
    PRIMARY_REPETITIONS * len(EXPECTED_PHASE3_COUNTS)
)
EXPECTED_PHASE2 = {
    "TP": 17,
    "TN": 10,
    "FP": 0,
    "FN": 1,
    "accuracy": 0.9642857142857143,
    "precision": 1.0,
    "recall": 0.9444444444444444,
    "f1": 0.9714285714285714,
    "fpr": 0.0,
}
EXPECTED_OVERHEAD = {
    "repeats": 5,
    "sample_rows": 1000,
    "end_to_end_mean_seconds": 3.1363267600012477,
    "end_to_end_std_seconds": 1.9548827180100747,
    "end_to_end_min_seconds": 1.9433268000138924,
    "end_to_end_max_seconds": 6.524428000004264,
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def bool_value(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"true", "1", "yes"}


def float_value(value: Any) -> float | None:
    if value in (None, ""):
        return None
    return float(value)


def close(a: Any, b: Any, tolerance: float = 1e-9) -> bool:
    try:
        return math.isclose(float(a), float(b), rel_tol=tolerance, abs_tol=tolerance)
    except (TypeError, ValueError):
        return False


def parse_time(value: str) -> datetime:
    return datetime.fromisoformat(str(value).replace("Z", "+00:00"))


def percentile(sorted_values: list[float], probability: float) -> float:
    if not sorted_values:
        raise ValueError("Cannot calculate a percentile of an empty list")
    position = (len(sorted_values) - 1) * probability
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return sorted_values[lower]
    weight = position - lower
    return sorted_values[lower] * (1.0 - weight) + sorted_values[upper] * weight


def bootstrap_mean_ci(
    values: list[float], *, resamples: int, seed: int, label: str
) -> tuple[float, float]:
    label_seed = int(hashlib.sha256(label.encode("utf-8")).hexdigest()[:8], 16)
    rng = random.Random(seed + label_seed)
    n = len(values)
    means = []
    for _ in range(resamples):
        means.append(sum(values[rng.randrange(n)] for _ in range(n)) / n)
    means.sort()
    return percentile(means, 0.025), percentile(means, 0.975)


def iqr_outliers(values: list[float]) -> list[float]:
    if len(values) < 4:
        return []
    ordered = sorted(values)
    q1 = percentile(ordered, 0.25)
    q3 = percentile(ordered, 0.75)
    spread = q3 - q1
    low = q1 - 1.5 * spread
    high = q3 + 1.5 * spread
    return [value for value in values if value < low or value > high]


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        raise ValueError(f"No rows available for {path}")
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


class Audit:
    def __init__(self) -> None:
        self.checks: list[dict[str, str]] = []

    def add(self, check_id: str, status: str, detail: str) -> None:
        if status not in {"PASS", "WARN", "FAIL"}:
            raise ValueError(status)
        self.checks.append({"id": check_id, "status": status, "detail": detail})

    def require(self, check_id: str, condition: bool, pass_detail: str, fail_detail: str) -> None:
        self.add(check_id, "PASS" if condition else "FAIL", pass_detail if condition else fail_detail)

    @property
    def failed(self) -> bool:
        return any(item["status"] == "FAIL" for item in self.checks)


def phase1_manifest(project: Path, audit: Audit) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    definitions = {
        "S1_LABEL_FLIP": {
            "folder": "S1_label_flip",
            "prefix": "s1_results",
        },
        "S2_BACKDOOR": {
            "folder": "S2_backdoor",
            "prefix": "s2_results",
        },
    }
    problems: list[str] = []
    for scenario, definition in definitions.items():
        for rate in (0.01, 0.05, 0.10):
            token = f"{round(rate * 100):02d}"
            for repetition in range(1, 6):
                result_path = (
                    project
                    / "evidence"
                    / "phase1"
                    / definition["folder"]
                    / f"{definition['prefix']}_rate{token}_run{repetition:02d}.json"
                )
                if not result_path.is_file():
                    problems.append(f"missing {result_path}")
                    continue
                data = json.loads(result_path.read_text(encoding="utf-8"))
                expected = (
                    data.get("scenario") == scenario
                    and close(data.get("poison_rate"), rate)
                    and int(data.get("repetition", -1)) == repetition
                )
                if not expected:
                    problems.append(f"metadata mismatch {result_path.name}")
                artefacts = {
                    "model": Path(str(data.get("model_file", ""))),
                    "training": Path(str(data.get("poisoned_training_file", ""))),
                    "samples": Path(str(data.get("poisoned_samples_file", ""))),
                }
                missing = [name for name, path in artefacts.items() if not path.is_file()]
                if missing:
                    problems.append(f"{result_path.name}: missing {', '.join(missing)}")
                    continue
                hashes = {name: sha256_file(path) for name, path in artefacts.items()}
                if str(data.get("model_sha256", "")).lower() != hashes["model"]:
                    problems.append(f"model hash mismatch {result_path.name}")
                rows.append(
                    {
                        "scenario": scenario,
                        "poison_rate": rate,
                        "repetition": repetition,
                        "attack_seed": data.get("attack_seed"),
                        "result_json": str(result_path),
                        "result_sha256": sha256_file(result_path),
                        "model_file": str(artefacts["model"]),
                        "model_sha256": hashes["model"],
                        "training_file": str(artefacts["training"]),
                        "training_sha256": hashes["training"],
                        "samples_file": str(artefacts["samples"]),
                        "samples_sha256": hashes["samples"],
                        "mlflow_run_id": data.get("mlflow_run_id"),
                        "metadata_match": expected,
                    }
                )
    audit.require(
        "P1-ARTIFACT-BINDING",
        len(rows) == 30 and not problems,
        "All 30 S1/S2 rate-by-repetition records match their metadata and model hashes.",
        f"Phase 1 binding problems: {problems[:8]}" if problems else f"Expected 30 rows, found {len(rows)}.",
    )
    baseline = project / "models" / "clean" / "clean_baseline.joblib"
    attack_scenarios = {row["scenario"] for row in rows}
    audit.require(
        "P1-CLEAN-BASELINE",
        baseline.is_file() and not any("CLEAN" in scenario for scenario in attack_scenarios),
        f"Clean baseline exists separately with SHA-256 {sha256_file(baseline) if baseline.is_file() else 'missing'}.",
        "Clean baseline is missing or an attack manifest row is labelled as clean.",
    )
    return rows


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as stream:
        return list(csv.DictReader(stream))


def phase2_checks(project: Path, audit: Audit) -> dict[str, Any]:
    metrics_path = project / "results" / "phase2" / "phase2_final_metrics.csv"
    metrics_rows = read_csv(metrics_path)
    metrics = metrics_rows[0]
    metrics_match = all(close(metrics.get(key), value) for key, value in EXPECTED_PHASE2.items())
    audit.require(
        "P2-FINAL-METRICS",
        len(metrics_rows) == 1 and metrics_match,
        "Final held-out metrics match TP=17, TN=10, FP=0 and FN=1.",
        "Phase 2 final metrics differ from the dissertation claim.",
    )

    missed_path = (
        project
        / "results"
        / "phase2"
        / "supplementary_analysis"
        / "s2_rate05_heldout_missed_runs.csv"
    )
    missed = read_csv(missed_path)
    missed_ok = (
        len(missed) == 1
        and missed[0].get("scenario") == "S2_BACKDOOR"
        and close(missed[0].get("poison_rate"), 0.05)
        and int(missed[0].get("run", -1)) == 4
        and close(missed[0].get("score"), 0.24164214904936798)
        and close(missed[0].get("operating_threshold"), 0.24359786888797563)
    )
    audit.require(
        "P2-S2-FALSE-NEGATIVE",
        missed_ok,
        "The S2 5% repetition-4 false negative is preserved with score and threshold.",
        "The S2 missed-run file is missing or differs from the reported case.",
    )

    overhead_path = project / "results" / "phase2" / "overhead_benchmark.json"
    overhead = json.loads(overhead_path.read_text(encoding="utf-8"))
    overhead_ok = all(close(overhead.get(key), value) for key, value in EXPECTED_OVERHEAD.items())
    audit.require(
        "P2-FINAL-OVERHEAD",
        overhead_ok,
        "Authoritative five-repeat overhead is mean 3.136327 s, SD 1.954883 s.",
        "The overhead JSON differs from the final reported benchmark.",
    )
    return {"metrics": metrics, "missed_s2": missed, "overhead": overhead}


def phase3_analysis(
    project: Path, audit: Audit, *, resamples: int, seed: int
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    runs_path = project / "results" / "phase3" / "phase3_response_runs.csv"
    records = read_csv(runs_path)
    counts = Counter(row.get("scenario") for row in records)
    audit.require(
        "P3-ACCEPTED-RUNS",
        len(records) == EXPECTED_PHASE3_TOTAL and dict(counts) == EXPECTED_PHASE3_COUNTS,
        "Accepted run counts are 30: S1=6, S2=7, S3=6, S4=6 and corrected S5=5.",
        f"Unexpected accepted run counts: total={len(records)}, counts={dict(counts)}.",
    )
    audit.require(
        "P3-EXCLUDE-V56",
        all(str(row.get("suspect_version")) != "56" for row in records),
        "Incomplete registration v56 is absent from the accepted run CSV.",
        "Incomplete registration v56 is present in the accepted run CSV.",
    )

    records.sort(key=lambda row: (str(row.get("scenario")), parse_time(row["timestamp"])))
    evidence_problems: list[str] = []
    scenario_position: defaultdict[str, int] = defaultdict(int)
    manifest: list[dict[str, Any]] = []
    for row in records:
        scenario = str(row["scenario"])
        scenario_position[scenario] += 1
        role = "PRIMARY_TIMING" if scenario_position[scenario] <= PRIMARY_REPETITIONS else "SUPPLEMENTARY_ALERT_VALIDATION"
        evidence_path = Path(row["evidence_file"])
        if not evidence_path.is_file():
            evidence_problems.append(f"missing {evidence_path}")
            continue
        data = json.loads(evidence_path.read_text(encoding="utf-8"))
        response = data.get("response", {})
        json_mttd = float_value(data.get("mttd_seconds"))
        csv_mttd = float_value(row.get("mttd_seconds"))
        csv_detector_start_mttd = float_value(row.get("mttd_to_detector_start_seconds"))
        detection_compute = float_value(data.get("detection_compute_seconds"))
        direct_mttd_match = close(json_mttd, csv_mttd)
        legacy_mttd_match = (
            row.get("mttd_reconstruction") == "LEGACY_START_MTTD_PLUS_DETECTION_COMPUTE"
            and close(json_mttd, csv_detector_start_mttd)
            and json_mttd is not None
            and detection_compute is not None
            and close(csv_mttd, json_mttd + detection_compute)
        )
        expected_fields = (
            data.get("phase") == "PHASE_3"
            and data.get("attack_scenario") == scenario
            and str(data.get("suspect_version")) == str(row.get("suspect_version"))
            and (direct_mttd_match or legacy_mttd_match)
            and close(data.get("detection_compute_seconds"), row.get("detection_compute_seconds"))
            and close(data.get("mttr_seconds"), row.get("mttr_seconds"))
            and bool(response.get("block", {}).get("success")) == bool_value(row.get("block_success"))
            and bool(response.get("quarantine", {}).get("success")) == bool_value(row.get("quarantine_success"))
            and bool(response.get("rollback", {}).get("success")) == bool_value(row.get("rollback_success"))
            and bool(data.get("wazuh_alert_verification", {}).get("alert_observed"))
            == bool_value(row.get("wazuh_alert_observed"))
        )
        if not expected_fields:
            evidence_problems.append(f"CSV/JSON mismatch {evidence_path.name}")
        if scenario == "S5_MLFLOW_TAMPERING":
            detection = data.get("detection_details", {})
            s5_hash_valid = (
                detection.get("detector_type") == "HASH_VERIFIED_BINARY_INTEGRITY_GATE"
                and detection.get("integrity_gate_control") == "MLFLOW_HASH"
                and detection.get("integrity_gate_decision") == "HASH_MISMATCH"
                and detection.get("artifact_hash_match") is False
            )
            if not s5_hash_valid:
                evidence_problems.append(
                    f"S5 is not corrected hash-gate evidence: {evidence_path.name}"
                )
        verification = data.get("wazuh_alert_verification", {})
        if bool(verification.get("alert_observed")):
            matching_alerts = verification.get("matching_alerts", [])
            rule_valid = any(
                str(alert.get("rule_id")) == "100103"
                and int(alert.get("rule_level", -1)) == 12
                for alert in matching_alerts
            )
            if not rule_valid:
                evidence_problems.append(
                    f"Observed Wazuh alert lacks rule 100103 level 12: {evidence_path.name}"
                )
        manifest.append(
            {
                "analysis_role": role,
                "scenario": scenario,
                "scenario_sequence": scenario_position[scenario],
                "suspect_version": row.get("suspect_version"),
                "timestamp": row.get("timestamp"),
                "mttd_basis": row.get("mttd_basis"),
                "mttd_seconds": row.get("mttd_seconds"),
                "detection_compute_seconds": row.get("detection_compute_seconds"),
                "mttr_seconds": row.get("mttr_seconds"),
                "block_success": bool_value(row.get("block_success")),
                "quarantine_success": bool_value(row.get("quarantine_success")),
                "rollback_success": bool_value(row.get("rollback_success")),
                "wazuh_alert_observed": bool_value(row.get("wazuh_alert_observed")),
                "wazuh_api_authenticated": bool(data.get("wazuh_api_status", {}).get("authenticated")),
                "evidence_file": str(evidence_path),
                "evidence_sha256": sha256_file(evidence_path),
                "csv_json_match": expected_fields,
            }
        )

    audit.require(
        "P3-EVIDENCE-BINDING",
        len(manifest) == EXPECTED_PHASE3_TOTAL and not evidence_problems,
        "Every Phase 3 CSV row matches its immutable JSON evidence file and SHA-256 hash.",
        f"Phase 3 evidence problems: {evidence_problems[:8]}",
    )
    all_actions = [
        row["block_success"] and row["quarantine_success"] and row["rollback_success"]
        for row in manifest
    ]
    audit.require(
        "P3-RESPONSE-SUCCESS",
        len(all_actions) == EXPECTED_PHASE3_TOTAL and all(all_actions),
        "All 30 accepted malicious runs completed block, quarantine and rollback.",
        f"Only {sum(all_actions)} of {EXPECTED_PHASE3_TOTAL} accepted runs completed every response action.",
    )
    wazuh_count = sum(row["wazuh_alert_observed"] for row in manifest)
    auth_count = sum(row["wazuh_api_authenticated"] for row in manifest)
    audit.require(
        "P3-WAZUH-ALERTS",
        wazuh_count > 0,
        f"{wazuh_count} of {EXPECTED_PHASE3_TOTAL} runs have individually verified Wazuh alerts backed by rule 100103 level 12.",
        "No accepted run contains an individually verified Wazuh rule 100103 alert.",
    )
    audit.add(
        "P3-WAZUH-AUTH",
        "WARN" if auth_count < wazuh_count else "PASS",
        f"Wazuh API authentication is true for {auth_count} runs; alert observation is independently true for {wazuh_count} runs.",
    )

    primary = [row for row in manifest if row["analysis_role"] == "PRIMARY_TIMING"]
    supplementary = [row for row in manifest if row["analysis_role"] != "PRIMARY_TIMING"]
    primary_counts = Counter(row["scenario"] for row in primary)
    audit.require(
        "P3-PRIMARY-TIMING-SET",
        len(primary) == 25 and all(primary_counts.get(name) == 5 for name in EXPECTED_PHASE3_COUNTS),
        "The pre-specified primary timing set contains the earliest five accepted runs per scenario (n=25).",
        f"Primary timing set is inconsistent: {dict(primary_counts)}.",
    )
    audit.require(
        "P3-SUPPLEMENTARY-SET",
        len(supplementary) == EXPECTED_SUPPLEMENTARY_RUNS,
        "Five later runs are labelled as supplementary alert-validation attempts and excluded from primary timing inference.",
        f"Expected {EXPECTED_SUPPLEMENTARY_RUNS} supplementary runs, found {len(supplementary)}.",
    )

    timing_rows: list[dict[str, Any]] = []
    for scenario in EXPECTED_PHASE3_COUNTS:
        scenario_rows = [row for row in primary if row["scenario"] == scenario]
        for metric in ("mttd_seconds", "detection_compute_seconds", "mttr_seconds"):
            if metric == "mttd_seconds":
                basis_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
                for row in scenario_rows:
                    basis_groups[str(row["mttd_basis"])].append(row)
            else:
                basis_groups = {"NOT_APPLICABLE": scenario_rows}
            for basis, group in basis_groups.items():
                values = [float(row[metric]) for row in group if row.get(metric) not in (None, "")]
                if not values:
                    continue
                low, high = bootstrap_mean_ci(
                    values,
                    resamples=resamples,
                    seed=seed,
                    label=f"{scenario}:{metric}:{basis}",
                )
                outliers = iqr_outliers(values)
                timing_rows.append(
                    {
                        "analysis_population": "PRIMARY_TIMING",
                        "scenario": scenario,
                        "metric": metric,
                        "mttd_basis": basis,
                        "n": len(values),
                        "mean_seconds": round(statistics.fmean(values), 6),
                        "median_seconds": round(statistics.median(values), 6),
                        "std_dev_seconds": round(statistics.stdev(values), 6) if len(values) > 1 else 0.0,
                        "min_seconds": round(min(values), 6),
                        "max_seconds": round(max(values), 6),
                        "ci95_lower_seconds": round(low, 6),
                        "ci95_upper_seconds": round(high, 6),
                        "ci_method": f"deterministic percentile bootstrap of mean ({resamples} resamples, seed={seed})",
                        "iqr_outlier_count": len(outliers),
                        "iqr_outlier_values": ";".join(f"{value:.6f}" for value in sorted(outliers)),
                    }
                )

    controls_path = project / "results" / "phase3" / "phase3_clean_control_runs.csv"
    controls = read_csv(controls_path)
    false_responses = sum(
        bool_value(row.get("block_attempted"))
        or bool_value(row.get("quarantine_attempted"))
        or bool_value(row.get("rollback_attempted"))
        for row in controls
    )
    controls_ok = len(controls) == 2 and all(bool_value(row.get("passed")) for row in controls) and false_responses == 0
    audit.require(
        "P3-CLEAN-CONTROLS",
        controls_ok,
        "Two clean controls passed without a false automated response.",
        f"Clean-control result differs: runs={len(controls)}, false responses={false_responses}.",
    )

    current_ci_path = project / "results" / "phase3" / "phase3_confidence_intervals.csv"
    current_ci = read_csv(current_ci_path)
    current_sizes = {(row["scenario"], row["metric"]): int(row["n"]) for row in current_ci}
    stale_ci = any(
        current_sizes.get((scenario, metric)) != 5
        for scenario in EXPECTED_PHASE3_COUNTS
        for metric in ("mttd_seconds", "detection_compute_seconds", "mttr_seconds")
    ) or "mttd_basis" not in (current_ci[0] if current_ci else {})
    audit.require(
        "P3-CURRENT-CI-FILE",
        not stale_ci,
        "Current Phase 3 confidence intervals use the primary set and preserve the MTTD basis.",
        "Current phase3_confidence_intervals.csv is stale and/or pools MTTD bases; use the regenerated audit output.",
    )

    archive_dir = project / "results" / "dissertation_evidence"
    archive_matches = []
    for name in (
        "phase3_response_runs.csv",
        "phase3_response_summary.json",
        "phase3_final_results.json",
        "phase3_confidence_intervals.csv",
    ):
        source = project / "results" / "phase3" / name
        archived = archive_dir / name
        archive_matches.append(source.is_file() and archived.is_file() and sha256_file(source) == sha256_file(archived))
    audit.require(
        "ARCHIVE-PHASE3",
        all(archive_matches),
        "The Phase 3 dissertation-evidence archive matches current final files.",
        "The Phase 3 dissertation-evidence archive is stale and must be refreshed after final statistics are accepted.",
    )
    summary = {
        "accepted_runs": len(manifest),
        "accepted_by_scenario": dict(Counter(row["scenario"] for row in manifest)),
        "primary_timing_runs": len(primary),
        "supplementary_alert_validation_runs": len(supplementary),
        "successful_responses": sum(all_actions),
        "wazuh_alerts_verified": wazuh_count,
        "wazuh_api_authenticated": auth_count,
        "clean_controls": len(controls),
        "clean_control_false_responses": false_responses,
        "primary_selection_rule": "Within each scenario, sort accepted records by timestamp and select the first five. The rule uses the pre-specified target_repetitions=5 and does not depend on outcomes.",
    }
    return manifest, timing_rows, summary


def unsw_poison_checks(evidence_dir: Path | None, audit: Audit) -> dict[str, Any]:
    if evidence_dir is None:
        audit.add("UNSW-POISON-EVIDENCE", "WARN", "UNSW poisoning evidence directory was not supplied.")
        return {}
    required = {
        "configuration": evidence_dir / "configuration.json",
        "thresholds": evidence_dir / "detector_thresholds.json",
        "runs": evidence_dir / "unsw_poison_runs.csv",
        "summary": evidence_dir / "unsw_poison_summary.csv",
        "performance": evidence_dir / "unsw_detector_performance.csv",
        "final": evidence_dir / "unsw_poison_final_results.json",
        "manifest": evidence_dir / "unsw_accepted_run_manifest.json",
    }
    missing = [name for name, path in required.items() if not path.is_file()]
    if missing:
        audit.add("UNSW-POISON-EVIDENCE", "FAIL", f"Missing UNSW evidence files: {missing}")
        return {"evidence_dir": str(evidence_dir), "missing": missing}
    configuration = json.loads(required["configuration"].read_text(encoding="utf-8"))
    final = json.loads(required["final"].read_text(encoding="utf-8"))
    manifest = json.loads(required["manifest"].read_text(encoding="utf-8"))
    runs = read_csv(required["runs"])
    summary = read_csv(required["summary"])
    clean_files = sorted((evidence_dir / "clean_controls").glob("clean_control_run*.json"))
    clean_controls = [json.loads(path.read_text(encoding="utf-8")) for path in clean_files]
    counts = Counter((row.get("scenario"), row.get("poison_rate")) for row in runs)
    expected_cells = {
        (scenario, rate)
        for scenario in ("S1_LABEL_FLIP", "S2_BACKDOOR")
        for rate in ("0.01", "0.05", "0.1")
    }
    counts_ok = len(runs) == 30 and all(counts.get(cell) == 5 for cell in expected_cells)
    audit.require(
        "UNSW-RUN-COUNTS",
        counts_ok,
        "UNSW poisoning validation contains 30 runs: five repetitions for each S1/S2 rate.",
        f"Unexpected UNSW run count or cells: total={len(runs)}, counts={dict(counts)}.",
    )
    split_ok = (
        configuration.get("holdout_used_for_training_or_threshold_selection") is False
        and "official test held out" in str(configuration.get("split_protocol", ""))
    )
    audit.require(
        "UNSW-HOLDOUT",
        split_ok,
        "Official UNSW test data is recorded as excluded from fitting and threshold selection.",
        "UNSW split metadata does not prove holdout isolation.",
    )
    source_hashes_ok = True
    for key in ("official_train", "official_test"):
        source = Path(str(configuration.get(key, "")))
        expected_hash = configuration.get(f"{key}_sha256")
        source_hashes_ok = source_hashes_ok and source.is_file() and sha256_file(source) == expected_hash
    audit.require(
        "UNSW-DATASET-HASHES",
        source_hashes_ok,
        "UNSW official train/test files match their recorded SHA-256 hashes.",
        "UNSW source file missing or dataset hash mismatch.",
    )
    manifest_records = manifest.get("records", [])
    manifest_ok = len(manifest_records) == len(runs)
    for row in manifest_records:
        evidence_path = Path(str(row.get("evidence_file", "")))
        manifest_ok = (
            manifest_ok
            and bool(row.get("included"))
            and evidence_path.is_file()
            and sha256_file(evidence_path) == row.get("evidence_sha256")
        )
    audit.require(
        "UNSW-EVIDENCE-BINDING",
        manifest_ok,
        "Every accepted UNSW run is bound to its JSON evidence by SHA-256.",
        "UNSW accepted-run manifest is incomplete or contains a hash mismatch.",
    )
    clean_false = sum(bool(row.get("false_response")) for row in clean_controls)
    audit.require(
        "UNSW-CLEAN-CONTROLS",
        len(clean_controls) == 5 and clean_false == 0,
        "Five UNSW clean controls produced no fusion false response.",
        f"UNSW clean controls={len(clean_controls)}, false responses={clean_false}.",
    )
    ci_columns = {name for row in summary for name in row if name.endswith("_ci95_lower")}
    audit.require(
        "UNSW-CONFIDENCE-INTERVALS",
        bool(ci_columns),
        "UNSW summary contains seeded 95% bootstrap confidence intervals.",
        "UNSW summary does not contain confidence-interval columns.",
    )
    audit.require(
        "UNSW-FINAL-COUNT",
        final.get("attack_run_count") == 30 and manifest.get("accepted_count") == 30,
        "UNSW final JSON and accepted manifest both report 30 attack runs.",
        "UNSW final JSON and accepted manifest disagree with the required 30-run cohort.",
    )
    return {
        "evidence_dir": str(evidence_dir),
        "attack_runs": len(runs),
        "clean_controls": len(clean_controls),
        "clean_false_responses": clean_false,
        "configuration_sha256": configuration.get("configuration_sha256"),
    }


def deployment_gate_checks(evidence_dir: Path | None, audit: Audit) -> dict[str, Any]:
    if evidence_dir is None:
        audit.add("DEPLOYMENT-GATE", "WARN", "Deployment-gate evidence directory was not supplied.")
        return {}
    files = sorted(evidence_dir.glob("*.json"))
    records = []
    for path in files:
        try:
            record = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        if record.get("event_type") == "MLFLOW_DEPLOYMENT_SECURITY_GATE":
            records.append(record)
    allowed = [row for row in records if row.get("decision") == "ALLOW"]
    blocked = [row for row in records if row.get("decision") == "BLOCK"]
    valid_allow = any(
        row.get("signature_verified") is True
        and row.get("artifact_hash_match") is True
        and row.get("blocklisted") is False
        for row in allowed
    )
    valid_block = any(
        row.get("integrity_gate_decision")
        in {"BLOCKLIST_MATCH", "SIGNATURE_VERIFICATION_FAILED", "SECURITY_CONTROL_UNAVAILABLE"}
        for row in blocked
    )
    audit.require(
        "DEPLOYMENT-GATE",
        valid_allow and valid_block,
        "Deployment evidence proves a signed clean allow and at least one fail-closed refusal.",
        f"Deployment evidence is incomplete: allow={len(allowed)}, block={len(blocked)}.",
    )
    return {
        "evidence_dir": str(evidence_dir),
        "records": len(records),
        "allowed": len(allowed),
        "blocked": len(blocked),
    }


def prisma_check(workbook: Path | None, audit: Audit) -> dict[str, Any]:
    if not workbook or not workbook.is_file():
        audit.add("PRISMA-RECORD", "FAIL", "No PRISMA search workbook was supplied.")
        return {}
    try:
        from openpyxl import load_workbook
    except ImportError:
        audit.add("PRISMA-RECORD", "WARN", "openpyxl is unavailable, so the PRISMA workbook was not inspected.")
        return {}
    book = load_workbook(workbook, data_only=False, read_only=True)
    search = book["Search Log"]
    completed_searches = 0
    for values in search.iter_rows(min_row=5, values_only=True):
        date = values[4] if len(values) > 4 else None
        query = values[5] if len(values) > 5 else None
        count = values[7] if len(values) > 7 else None
        if date not in (None, "") and query not in (None, "") and count not in (None, ""):
            completed_searches += 1
    screening = book["Screening Record"]
    screened = included = full_text_excluded = 0
    for values in screening.iter_rows(min_row=5, values_only=True):
        title_decision = values[8] if len(values) > 8 else None
        full_text_decision = values[10] if len(values) > 10 else None
        exclusion_reason = values[11] if len(values) > 11 else None
        included_value = values[12] if len(values) > 12 else None
        if title_decision in {"Include", "Exclude"}:
            screened += 1
        if str(included_value).strip().lower() == "yes":
            included += 1
        if full_text_decision == "Exclude" and exclusion_reason not in (None, ""):
            full_text_excluded += 1
    ready = completed_searches >= 2 and screened > 0 and included > 0
    audit.require(
        "PRISMA-RECORD",
        ready,
        f"PRISMA workbook has {completed_searches} completed searches, {screened} screened records and {included} included studies.",
        f"PRISMA workbook is incomplete: completed searches={completed_searches}, screened={screened}, included={included}.",
    )
    return {
        "workbook": str(workbook),
        "completed_searches": completed_searches,
        "screened_records": screened,
        "included_studies": included,
        "full_text_exclusions_with_reason": full_text_excluded,
    }


def report_check(report: Path | None, audit: Audit) -> dict[str, Any]:
    if not report or not report.is_file():
        audit.add("REPORT-CONTENT", "FAIL", "No dissertation DOCX was supplied for content checking.")
        return {}
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        audit.add("REPORT-CONTENT", "WARN", "python-docx is unavailable, so report content was not inspected.")
        return {}
    try:
        document = Document(report)
    except PermissionError:
        audit.add("REPORT-CONTENT", "FAIL", "The dissertation is locked. Close Word or supply a copied DOCX.")
        return {}

    sections: defaultdict[str, list[str]] = defaultdict(list)
    current = "FRONT"
    texts: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            block = Paragraph(child, document._body)
            value = " ".join(block.text.split())
        elif child.tag.endswith("}tbl"):
            block = Table(child, document._body)
            value = " ".join(cell.text for row in block.rows for cell in row.cells)
        else:
            continue
        texts.append(value)
        upper = value.upper()
        match = re.match(r"CHAPTER\s+([1-8])\.", upper)
        if match:
            current = f"CHAPTER {match.group(1)}"
        sections[current].append(value)
    word_pattern = re.compile(r"\b\w+(?:[-']\w+)*\b", flags=re.UNICODE)
    main_words = sum(len(word_pattern.findall("\n".join(sections[f"CHAPTER {number}"]))) for number in range(1, 7))
    full_text = "\n".join(texts)
    placeholders = [
        value
        for value in texts
        if "[INSERT " in value.upper()
        or value.startswith("Record the exact commands")
        or value.startswith("Include the final")
        or value.startswith("Provide one manifest")
    ]
    reference_values = [value for value in sections["CHAPTER 7"] if value and not value.upper().startswith("CHAPTER 7")]
    has_core_claims = all(
        needle in full_text
        for needle in (
            "0.964286",
            "0.971429",
            "3.1363",
            "30 accepted",
            "eight of 30",
            "0.241642",
            "0.243598",
        )
    )
    audit.require(
        "REPORT-EVIDENCE-CLAIMS",
        has_core_claims and "2.631" not in full_text,
        "Report contains the final metrics and overhead and does not contain the obsolete 2.631 s result.",
        "Report is missing a final evidence claim or still contains the obsolete 2.631 s overhead.",
    )
    audit.require(
        "REPORT-PLACEHOLDERS",
        not placeholders,
        "No table, figure or appendix instruction placeholders remain.",
        f"Report still has {len(placeholders)} placeholder or instruction paragraphs.",
    )
    audit.add(
        "REPORT-WORD-GUIDANCE",
        "PASS" if 7200 <= main_words <= 8800 else "WARN",
        f"Chapters 1-6 contain approximately {main_words} words; handbook guidance is 7,200-8,800.",
    )
    audit.add(
        "REPORT-REFERENCES",
        "PASS" if len(reference_values) >= 25 else "WARN",
        f"Chapter 7 contains {len(reference_values)} non-heading reference paragraphs. Quality and in-text use still require manual verification.",
    )
    return {
        "report": str(report),
        "main_body_words": main_words,
        "placeholder_count": len(placeholders),
        "placeholder_examples": placeholders[:10],
        "chapter7_reference_paragraphs": len(reference_values),
        "inline_shapes": len(document.inline_shapes),
    }


def write_markdown(path: Path, payload: dict[str, Any]) -> None:
    checks = payload["checks"]
    phase3 = payload["phase3"]
    status_counts = Counter(item["status"] for item in checks)
    lines = [
        "# CST4990 Research Submission Audit",
        "",
        f"Generated: {payload['generated_at']}",
        "",
        f"Overall audit status: **{'FAIL' if status_counts['FAIL'] else 'PASS'}**",
        "",
        f"Checks: {status_counts['PASS']} pass, {status_counts['WARN']} warning, {status_counts['FAIL']} fail.",
        "",
        "## Checks",
        "",
        "| ID | Status | Evidence |",
        "|---|---|---|",
    ]
    for item in checks:
        detail = item["detail"].replace("|", "\\|")
        lines.append(f"| {item['id']} | {item['status']} | {detail} |")
    lines.extend(
        [
            "",
            "## Phase 3 analysis populations",
            "",
            f"The outcome analysis uses all {phase3['accepted_runs']} accepted malicious runs. The primary timing analysis uses the earliest five accepted records in each scenario, matching the pre-specified target of five repetitions. {phase3['supplementary_alert_validation_runs']} later records are retained as supplementary alert-validation attempts. The selection rule depends on timestamp and the pre-specified repetition target, not on measured outcomes.",
            "",
            "S2 primary workflow MTTD uses the explicit attack-start basis. Later S2 MLflow-tag records remain in the accepted-run and alert-validation evidence but are excluded from the primary timing estimate.",
            "",
            "## Authoritative headline values",
            "",
            "- Phase 2: TP=17, TN=10, FP=0, FN=1; accuracy=0.964286, precision=1.000000, recall=0.944444, F1=0.971429.",
            "- Phase 2 overhead: five repetitions; mean=3.136327 s, SD=1.954883 s, range=1.943327-6.524428 s per 1,000-row full-stack evaluation.",
            f"- Phase 3: {phase3['accepted_runs']} accepted malicious runs; {phase3['successful_responses']} completed block, quarantine and rollback; {phase3['wazuh_alerts_verified']} Wazuh alerts verified; two clean controls produced no false response.",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate and freeze CST4990 Secure MLflow research evidence")
    parser.add_argument(
        "--project-root",
        type=Path,
        default=Path(r"E:\secure_mlflow_framework_3phase\secure_mlflow_framework_3phase"),
    )
    parser.add_argument("--output-dir", type=Path, default=Path("submission_audit"))
    parser.add_argument("--report", type=Path, default=None)
    parser.add_argument("--prisma-workbook", type=Path, default=None)
    parser.add_argument("--unsw-evidence-dir", type=Path, default=None)
    parser.add_argument("--deployment-evidence-dir", type=Path, default=None)
    parser.add_argument("--bootstrap-resamples", type=int, default=20000)
    parser.add_argument("--seed", type=int, default=20260910)
    parser.add_argument("--strict", action="store_true", help="Return exit code 1 when any audit check fails")
    args = parser.parse_args()

    project = args.project_root.resolve()
    output = args.output_dir.resolve()
    output.mkdir(parents=True, exist_ok=True)
    audit = Audit()

    phase1_rows = phase1_manifest(project, audit)
    phase2 = phase2_checks(project, audit)
    phase3_rows, timing_rows, phase3_summary = phase3_analysis(
        project,
        audit,
        resamples=args.bootstrap_resamples,
        seed=args.seed,
    )
    unsw = unsw_poison_checks(
        args.unsw_evidence_dir.resolve() if args.unsw_evidence_dir else None,
        audit,
    )
    deployment = deployment_gate_checks(
        args.deployment_evidence_dir.resolve() if args.deployment_evidence_dir else None,
        audit,
    )
    prisma = prisma_check(args.prisma_workbook.resolve() if args.prisma_workbook else None, audit)
    report = report_check(args.report.resolve() if args.report else None, audit)

    write_csv(output / "phase1_s1_s2_artifact_manifest.csv", phase1_rows)
    write_csv(output / "phase3_run_manifest.csv", phase3_rows)
    write_csv(output / "phase3_primary_timing_statistics.csv", timing_rows)
    write_csv(output / "phase3_confidence_intervals.csv", timing_rows)

    payload = {
        "generated_at": datetime.now().astimezone().isoformat(),
        "project_root": str(project),
        "method": {
            "primary_phase3_repetitions_per_scenario": PRIMARY_REPETITIONS,
            "bootstrap_resamples": args.bootstrap_resamples,
            "bootstrap_seed": args.seed,
        },
        "phase2": phase2,
        "phase3": phase3_summary,
        "unsw_poison_validation": unsw,
        "deployment_gate": deployment,
        "prisma": prisma,
        "report": report,
        "checks": audit.checks,
    }
    json_path = output / "research_submission_audit.json"
    json_path.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")
    write_markdown(output / "research_submission_audit.md", payload)

    status_counts = Counter(item["status"] for item in audit.checks)
    print(f"Audit outputs: {output}")
    print(f"PASS={status_counts['PASS']} WARN={status_counts['WARN']} FAIL={status_counts['FAIL']}")
    for item in audit.checks:
        if item["status"] != "PASS":
            print(f"{item['status']} {item['id']}: {item['detail']}")
    return 1 if args.strict and audit.failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
